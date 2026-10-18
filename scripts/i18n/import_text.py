"""Reads the owner's text workbook and updates the catalogs.

    python3 scripts/i18n/import_text.py tutorio-ui-text.xlsx

This only turns the "Text" and "_map" sheets into JSON. Every check (unknown, missing or repeated keys,
placeholders, spaces) and every write is done by `node scripts/i18n/messages.mjs apply`, which writes nothing
unless all rows are valid. Only the `en` and `vi` columns are imported for now.
"""

import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]


def cell(value):
    """Text and empty cells as they are; anything else is passed on so the check can refuse it."""
    if value is None or isinstance(value, str):
        return value
    return {"not text": str(value)}


def read(path):
    wb = load_workbook(path, read_only=True)
    for name in ("Text", "_map"):
        if name not in wb.sheetnames:
            sys.exit(f"The workbook has no '{name}' sheet.")
    text = list(wb["Text"].iter_rows(values_only=True))
    if not text:
        sys.exit("The 'Text' sheet is empty.")
    header = [str(h).strip() if h is not None else "" for h in text[0]]
    columns = {}
    for name in ("key", "en", "vi"):
        if header.count(name) != 1:
            sys.exit(f"The 'Text' sheet needs exactly one '{name}' column.")
        columns[name] = header.index(name)

    source = {}
    for row in list(wb["_map"].iter_rows(values_only=True))[1:]:
        if row and isinstance(row[0], str) and len(row) > 1:
            source.setdefault(row[0], row[1])

    rows = []
    for values in text[1:]:
        if all(v is None for v in values):
            continue
        get = lambda name: values[columns[name]] if columns[name] < len(values) else None
        key = get("key")
        rows.append(
            {
                "key": cell(key),
                "source": cell(source.get(key)) if isinstance(key, str) else None,
                "en": cell(get("en")),
                "vi": cell(get("vi")),
            }
        )
    return rows


def main():
    if len(sys.argv) != 2:
        sys.exit("Usage: python3 scripts/i18n/import_text.py <file.xlsx>")
    rows = read(sys.argv[1])
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "rows.json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump(rows, f, ensure_ascii=False)
        result = subprocess.run(["node", str(ROOT / "scripts/i18n/messages.mjs"), "apply", path])
    sys.exit(result.returncode)


if __name__ == "__main__":
    main()
